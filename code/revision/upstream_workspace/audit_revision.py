"""Independent calculation-route and evidence-accounting checks, not truth validation."""
from pathlib import Path
import json,hashlib,sys,re,zipfile
import numpy as np,pandas as pd
import statsmodels.api as sm
from scipy.special import logit
from docx import Document
from docx.oxml.ns import qn
O=Path(__file__).resolve().parent;ROOT=O.parents[3];T=O/'tables';P=O/'author_package'
sys.path.insert(0,str(ROOT));from research_pipeline.common import fit_linear
protocol=json.loads((O/'FROZEN_REVISION_PROTOCOL.json').read_text())
checks=[]
def check(name,condition,detail):
 checks.append({'check':name,'passed':bool(condition),'detail':detail})
 if not condition:raise AssertionError((name,detail))
d=pd.read_csv(T/'annual_model_inputs_revision.csv')
check('Annual class and missing-area accounting',np.allclose(d[['nonwater_area_km2','seasonal_water_area_km2','permanent_water_area_km2']].sum(axis=1),d.observed_annual_area_km2,atol=1e-6) and np.allclose(d.observed_annual_area_km2+d.missing_annual_area_km2,d.reference_area_km2,atol=1e-6),'700 genuine cohort/region/year summaries; missing kept separate')
for cls in ['seasonal','permanent']:
 check(cls+' missing-support bounds',(d[f'{cls}_reference_upper_fraction']>=d[f'{cls}_reference_lower_fraction']).all() and np.allclose(d[f'{cls}_reference_upper_fraction']-d[f'{cls}_reference_lower_fraction'],1-d.coverage_fraction,atol=1e-10),'Bounds allocate only missing area, not physical classification error')
fits=0;maxerr=0
for (cohort,region),g in d[d.year.le(2021)&d.eligible].groupby(['cohort','region']):
 g=g.sort_values('year')
 if not np.all(np.diff(g.year)==1):continue
 for model,features in protocol['models'].items():
  for cls in ['seasonal','permanent']:
   y=logit(g[f'{cls}_fraction_of_observed'].clip(1e-6,1-1e-6));x=g[features]
   f=fit_linear(x,y,g.year,2)
   independent=sm.OLS(y,sm.add_constant(x,has_constant='add')).fit(cov_type='HAC',cov_kwds={'maxlags':2,'use_correction':True},use_t=True)
   err=max(np.max(np.abs(f['coefficients']-independent.params.to_numpy())),np.max(np.abs(f['standard_errors']-independent.bse.to_numpy())))
   maxerr=max(maxerr,float(err));fits+=1
check('Alternative statsmodels OLS/HAC implementation',maxerr<1e-6,{'fits_compared':fits,'max_absolute_coefficient_or_SE_difference':maxerr,'boundary':'separate implementation, same inputs; not independent environmental validation'})
b=pd.read_csv(T/'block_bootstrap_revision.csv');gap=b.interval_status.eq('UNAVAILABLE_calendar_gaps')
check('Calendar-gap bootstrap guard',gap.any() and b.loc[gap,['lower95','upper95']].isna().all().all() and b.loc[gap,'replicates'].eq(0).all() and b.scope.isin(['primary','2013_2021']).all(),'No compressed-gap intervals or scope-key overwrite')
v=pd.read_csv(T/'chronological_cv_revision.csv');out=[]
keys=['cohort','region','source_class']
for key,g in v[v.fold.eq('two_year')].groupby(keys):
 for model in protocol['models']:
  selected=g[g.model.eq(model)&g.status.eq('PASS')]
  if selected.empty:continue
  rec=dict(zip(keys,key));rec.update(model=model,accepted_folds=len(selected),test_years=int(selected.n_test.sum()))
  for method,name in [(model,'rmse'),('training_mean','mean_benchmark_rmse'),('trend','trend_benchmark_rmse')]:
   a=g[g.model.eq(method)&g.status.eq('PASS')&g.test_start.isin(selected.test_start)]
   check('CV matched fold '+str(key)+model+method,len(a)==len(selected),'Same selected test starts and sample counts')
   rec[name]=float(np.sqrt(a.squared_error_sum.sum()/a.n_test.sum()))
  out.append(rec)
pd.DataFrame(out).to_csv(T/'chronological_cv_matched_fold_summary.csv',index=False)
sim=pd.read_csv(T/'simulated_detectability_records.csv');check('Synthetic accounting',len(sim)==25920,'120 starts × 3 durations × 6 deficits × 12 region/cohort combinations; not actual droughts')
summary=pd.read_csv(T/'simulated_detectability_summary.csv');check('Synthetic positive-control failure retained',summary.detected_deficits.eq(0).all() and summary.injected_deficit_fraction.max()==.2,'No detection even at 0.20; no measured recovery or empirical specificity claimed')
c=pd.read_csv(T/'static_context_areas.csv');check('Context source records',len(c)>0 and (O/'CONTEXT_COMPLETE.json').exists(),'Actual downloaded HydroRIVERS and WorldCover sources, checksummed separately')
blocked=json.loads((O/'external_sources/context_acquisition_log.json').read_text())
check('Blocked reference download not used as labels',any(x.get('scientific_data_acquired') is False for x in blocked),'CyVerse HTML blocker recorded as failure, not acquired data')
old=ROOT/'Frontiers_Earth_Science_Inline_20261016/manuscript.docx'
check('Original manuscript unchanged',hashlib.sha256(old.read_bytes()).hexdigest()=='d6051f453964fb59336c80c2027f2020780ce2f1a3a48a8bcf9e039598ab6c9e','All edits in separate author_package')
doc=Document(P/'manuscript.docx');txt=lambda p:''.join(x.text or '' for x in p._p.iter(qn('w:t')))
body='\n'.join(txt(p) for p in doc.paragraphs);capt=[txt(p) for p in doc.paragraphs if p.style.name=='Caption']
check('Figure/table counts',len(doc.inline_shapes)==12 and len(doc.tables)==3 and len([x for x in capt if x.startswith('Figure ')])==12,'Twelve embedded figures and three real Word tables')
bookmarks={x.get(qn('w:name')) for x in doc._element.iter(qn('w:bookmarkStart'))}
targets=[]
for f in doc._element.iter(qn('w:fldSimple')):
 m=re.search(r'REF\s+(\S+)',f.get(qn('w:instr'),''))
 if m:targets.append(m[1])
check('Native Word cross-reference targets',all(x in bookmarks for x in targets),{'REF_fields':len(targets),'unresolved':sorted(set(targets)-bookmarks)})
check('No visible broken reference markers',all(s not in body for s in ['Error! Reference source not found','??','INLINEFIGURE','{{WORD_COUNT}}']),'Word field update also checked in rendered PDF')
refs=json.loads((P/'references/references.json').read_text())
dois=[r.get('DOI','').lower() for r in refs if r.get('DOI')]
check('Reference identities unique',len(refs)==22 and len(dois)==len(set(dois)),'Nineteen original plus three registry-verified additions')
bib=[txt(p) for p in doc.paragraphs if p.style.name=='Bibliography1'];check('Reference list count',len(bib)==22,'All original references retained')
check('New citations called in text',all(x in body.split('References')[0] for x in ['Lehner and Grill, 2013','Zanaga et al., 2021','Mukherjee et al., 2024']),'Static author-date calls, not invented live Zotero fields')
result={'assurance':'self_review_computational_and_document_checks_only','independent_accuracy_established':False,'actual_recovery_established':False,'passed':all(x['passed'] for x in checks),'checks':checks}
(O/'revision_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print('Passed',len(checks),'checks; no independent accuracy or actual recovery claimed.')
